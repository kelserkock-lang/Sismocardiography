"""
RELATÓRIO DE RESULTADOS - PIBIC SISMOCARDIOGRAFIA
==================================================
Gera um relatório Word (.docx) comparando os grupos
Saudável x Não saudável a partir dos dados de sismocardiografia
(cache de pacientes) e da planilha Pesquisa_PIBIC_sismo.xlsx.

Conteúdo:
  - Dados sociodemográficos (idade, sexo, n por grupo)
  - Frequência cardíaca (FC)
  - Variabilidade da FC (HRV: SDNN, RMSSD)
  - Frequência respiratória (FR) estimada
  - Morfologia média da onda SCG por grupo
  - Dashboards (onda média de cada paciente por grupo)
"""

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy import signal, stats
from pathlib import Path
import pickle
import warnings
warnings.filterwarnings('ignore')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DATA_DIR = Path(r"C:\Users\kelse\OneDrive\Área de Trabalho\DADOS_PIBIC")
CACHE_DIR = DATA_DIR / "_cache_pacientes"
FIG_DIR = DATA_DIR / "relatorio_figuras"
FS = 100.0

COR_SAUD = '#2ecc71'
COR_NAO = '#e74c3c'

# Pontos fiduciais (rótulo, cor) — Sørensen et al. 2018
PONTOS = ['Bs', 'Cs', 'Es', 'Gs', 'Ks', 'Bd', 'Fd', 'Gd']
FID_INFO = {
    'Bs': ('Bs', '#e6194B'), 'Cs': ('Cs', '#f58231'),
    'Es': ('Es/MC', '#ffbe0b'), 'Gs': ('Gs/AO', '#3cb44b'),
    'Ks': ('Ks', '#42d4f4'), 'Bd': ('Bd/AC', '#4363d8'),
    'Fd': ('Fd/MO', '#911eb4'), 'Gd': ('Gd', '#f032e6'),
}
# Intervalos temporais: (chave, rótulo, ponto_inicial, ponto_final)
INTERVALOS = [
    ('civ',  'CIV: MC→AO (ms)',        'Es', 'Gs'),
    ('lvet', 'LVET: AO→AC (ms)',       'Gs', 'Bd'),
    ('riv',  'RIV: AC→MO (ms)',        'Bd', 'Fd'),
    ('ench', 'Enchim.: MO→pico E (ms)', 'Fd', 'Gd'),
]

plt.rcParams['figure.dpi'] = 150
plt.rcParams['font.size'] = 10


# ============================================================
# UTILIDADES
# ============================================================

def is_saudavel(grupo):
    g = str(grupo).lower()
    return ('não' not in g) and ('nao' not in g) and ('saud' in g)


def calc_hrv(peaks, fs=FS):
    """Calcula FC, SDNN e RMSSD a partir dos picos (batimentos)."""
    if len(peaks) < 4:
        return np.nan, np.nan, np.nan
    rr = np.diff(np.sort(peaks)) / fs * 1000.0  # ms
    # remover intervalos fisiologicamente implausíveis
    rr = rr[(rr > 300) & (rr < 2000)]
    if len(rr) < 3:
        return np.nan, np.nan, np.nan
    fc = 60000.0 / np.mean(rr)
    sdnn = np.std(rr, ddof=1)
    rmssd = np.sqrt(np.mean(np.diff(rr) ** 2))
    return fc, sdnn, rmssd


def estimar_fr(sinal, fs=FS):
    """
    Estima a frequência respiratória (rpm) pela modulação de amplitude
    do SCG: envelope de Hilbert -> PSD na banda respiratória (0.1-0.5 Hz).
    """
    if len(sinal) < int(20 * fs):
        return np.nan
    env = np.abs(signal.hilbert(sinal))
    env = env - np.mean(env)
    # suavizar envelope (passa-baixa 1 Hz)
    b, a = signal.butter(4, 1.0 / (fs / 2), btype='low')
    env = signal.filtfilt(b, a, env)
    f, pxx = signal.welch(env, fs=fs, nperseg=min(len(env), int(30 * fs)))
    banda = (f >= 0.1) & (f <= 0.5)
    if not np.any(banda):
        return np.nan
    f_pico = f[banda][np.argmax(pxx[banda])]
    return f_pico * 60.0


def fmt_ms(x):
    return "—" if np.isnan(x) else f"{x:.1f}"


def resumo_grupo(valores):
    v = np.array([x for x in valores if not np.isnan(x)])
    if len(v) == 0:
        return "—"
    return f"{np.mean(v):.1f} ± {np.std(v, ddof=1):.1f}"


def teste_grupos(a, b):
    """Mann-Whitney U entre dois grupos; retorna p-valor formatado."""
    a = np.array([x for x in a if not np.isnan(x)])
    b = np.array([x for x in b if not np.isnan(x)])
    if len(a) < 3 or len(b) < 3:
        return "—"
    try:
        _, p = stats.mannwhitneyu(a, b, alternative='two-sided')
    except ValueError:
        return "—"
    if p < 0.001:
        return "< 0,001"
    return f"{p:.3f}".replace('.', ',')


def carregar_fiducial_times(num):
    """Carrega os tempos (s) dos pontos fiduciais automáticos do paciente."""
    p = DATA_DIR / f"anotacoes_auto_paciente_{num}.pkl"
    if not p.exists():
        return {}
    with open(p, 'rb') as f:
        d = pickle.load(f)
    return d.get('fiducial_times', {})


def calc_intervalos(ft):
    """Calcula os intervalos temporais (ms) a partir dos pontos fiduciais."""
    out = {}
    for chave, _lab, a, b in INTERVALOS:
        if a in ft and b in ft:
            out[chave] = (ft[b] - ft[a]) * 1000.0
        else:
            out[chave] = np.nan
    return out


# ============================================================
# CARGA DE DADOS
# ============================================================

def carregar_dados():
    df = pd.read_excel(DATA_DIR / "Pesquisa_PIBIC_sismo.xlsx")
    df.columns = ['Numero', 'Idade', 'Sexo', 'Grupo']
    df['Numero'] = df['Numero'].astype(int)
    df['Sexo'] = df['Sexo'].astype(str).str.strip().str.upper().str[0]
    df = df.drop_duplicates(subset=['Numero'])
    demo = {int(r.Numero): (int(r.Idade), r.Sexo, r.Grupo) for r in df.itertuples()}

    registros = []
    for num in range(1, 61):
        cache = CACHE_DIR / f"paciente_{num}.npz"
        if not cache.exists():
            continue
        d = np.load(cache)
        cycles = d['cycles']
        if len(cycles) < 2:
            continue
        idade, sexo, grupo = demo.get(num, (int(d['idade']), str(d['sexo']), str(d['grupo'])))
        fc, sdnn, rmssd = calc_hrv(d['peaks'])
        if np.isnan(fc):
            fc = float(d['fc_est'])
        fr = estimar_fr(d['sinal'])
        ensemble = np.mean(cycles, axis=0)
        ft = carregar_fiducial_times(num)
        registros.append({
            'num': num, 'idade': idade, 'sexo': sexo, 'grupo': grupo,
            'saudavel': is_saudavel(grupo),
            'fc': fc, 'sdnn': sdnn, 'rmssd': rmssd, 'fr': fr,
            'n_ciclos': len(cycles), 'ensemble': ensemble,
            **{f't_{k}': ft.get(k, np.nan) for k in PONTOS},
            **calc_intervalos(ft),
        })
    return pd.DataFrame(registros)


# ============================================================
# FIGURAS
# ============================================================

def fig_sociodemografico(df, path):
    saud = df[df['saudavel']]
    nao = df[~df['saudavel']]
    fig, axes = plt.subplots(1, 3, figsize=(13, 4))

    # (a) n por grupo
    axes[0].bar(['Saudável', 'Não saudável'], [len(saud), len(nao)],
                color=[COR_SAUD, COR_NAO], edgecolor='black')
    for i, v in enumerate([len(saud), len(nao)]):
        axes[0].text(i, v + 0.3, str(v), ha='center', fontweight='bold')
    axes[0].set_title('(a) Nº de participantes')
    axes[0].set_ylabel('Nº')

    # (b) sexo por grupo
    def cont_sexo(g):
        return (np.sum(g['sexo'] == 'M'), np.sum(g['sexo'] == 'F'))
    ms, fs_ = cont_sexo(saud)
    mn, fn = cont_sexo(nao)
    x = np.arange(2)
    axes[1].bar(x, [ms, mn], label='M', color='#3498db', edgecolor='black')
    axes[1].bar(x, [fs_, fn], bottom=[ms, mn], label='F', color='#e84393', edgecolor='black')
    axes[1].set_xticks(x)
    axes[1].set_xticklabels(['Saudável', 'Não saudável'])
    axes[1].set_title('(b) Distribuição por sexo')
    axes[1].set_ylabel('Nº')
    axes[1].legend()

    # (c) idade por grupo
    dados_idade = [saud['idade'].values, nao['idade'].values]
    bp = axes[2].boxplot(dados_idade, labels=['Saudável', 'Não saudável'],
                         patch_artist=True, widths=0.5)
    for patch, cor in zip(bp['boxes'], [COR_SAUD, COR_NAO]):
        patch.set_facecolor(cor)
        patch.set_alpha(0.6)
    axes[2].set_title('(c) Idade')
    axes[2].set_ylabel('Idade (anos)')

    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)


def fig_morfologia(df, path):
    ta = np.linspace(-0.3, 0.6, 90)

    def grand_avg(sub, normalizar=True):
        arr = []
        for e in sub['ensemble']:
            e = np.asarray(e)
            if len(e) != 90:
                continue
            if normalizar:
                m = np.max(np.abs(e))
                e = e / m if m > 0 else e
            arr.append(e)
        arr = np.array(arr)
        return np.mean(arr, axis=0), np.std(arr, axis=0)

    saud = df[df['saudavel']]
    nao = df[~df['saudavel']]

    fig, axes = plt.subplots(1, 2, figsize=(13, 4.5))

    # normalizado (forma)
    for sub, cor, lab in [(saud, COR_SAUD, 'Saudável'), (nao, COR_NAO, 'Não saudável')]:
        m, s = grand_avg(sub, normalizar=True)
        axes[0].plot(ta, m, color=cor, lw=2.3, label=f'{lab} (n={len(sub)})')
        axes[0].fill_between(ta, m - s, m + s, color=cor, alpha=0.15)
    axes[0].axvline(0, color='gray', ls='--', alpha=0.5)
    axes[0].set_title('(a) Morfologia média normalizada')
    axes[0].set_xlabel('Tempo (s)')
    axes[0].set_ylabel('Amplitude normalizada')
    axes[0].legend()
    axes[0].grid(alpha=0.3)

    # amplitude absoluta
    for sub, cor, lab in [(saud, COR_SAUD, 'Saudável'), (nao, COR_NAO, 'Não saudável')]:
        m, s = grand_avg(sub, normalizar=False)
        axes[1].plot(ta, m, color=cor, lw=2.3, label=f'{lab} (n={len(sub)})')
        axes[1].fill_between(ta, m - s, m + s, color=cor, alpha=0.15)
    axes[1].axvline(0, color='gray', ls='--', alpha=0.5)
    axes[1].set_title('(b) Amplitude média absoluta')
    axes[1].set_xlabel('Tempo (s)')
    axes[1].set_ylabel('Aceleração (m/s²)')
    axes[1].legend()
    axes[1].grid(alpha=0.3)

    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)


def fig_boxplots(df, path):
    saud = df[df['saudavel']]
    nao = df[~df['saudavel']]
    metricas = [('fc', 'FC (bpm)'), ('sdnn', 'SDNN (ms)'),
                ('rmssd', 'RMSSD (ms)'), ('fr', 'FR (rpm)')]
    fig, axes = plt.subplots(1, 4, figsize=(15, 4))
    for ax, (col, titulo) in zip(axes, metricas):
        ds = saud[col].dropna().values
        dn = nao[col].dropna().values
        bp = ax.boxplot([ds, dn], labels=['Saud.', 'Não\nsaud.'],
                        patch_artist=True, widths=0.5, showmeans=True)
        for patch, cor in zip(bp['boxes'], [COR_SAUD, COR_NAO]):
            patch.set_facecolor(cor)
            patch.set_alpha(0.6)
        ax.set_title(titulo, fontweight='bold')
        ax.grid(alpha=0.3, axis='y')
    fig.suptitle('Comparação de métricas fisiológicas entre grupos', fontweight='bold')
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)


def fig_dashboard(sub, titulo, cor, path, cols=5):
    ta = np.linspace(-0.3, 0.6, 90)
    n = len(sub)
    rows = max(1, (n + cols - 1) // cols)
    fig, axes = plt.subplots(rows, cols, figsize=(cols * 2.6, rows * 2.0))
    axes = np.array(axes).reshape(-1)
    for ax in axes:
        ax.axis('off')
    for i, (_, p) in enumerate(sub.iterrows()):
        ax = axes[i]
        ax.axis('on')
        e = np.asarray(p['ensemble'])
        ax.plot(ta, e, color=cor, lw=1.4)
        ax.axvline(0, color='gray', ls=':', alpha=0.5)
        ax.set_title(f"#{p['num']} | {p['sexo']} {p['idade']}a\nFC={p['fc']:.0f} FR={p['fr']:.0f}",
                     fontsize=7)
        ax.set_xticks([]); ax.set_yticks([])
    fig.suptitle(titulo, fontweight='bold', fontsize=13)
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)


def fig_intervalos(df, path):
    saud = df[df['saudavel']]
    nao = df[~df['saudavel']]
    fig, axes = plt.subplots(1, len(INTERVALOS), figsize=(15, 4))
    for ax, (chave, lab, _a, _b) in zip(axes, INTERVALOS):
        ds = saud[chave].dropna().values
        dn = nao[chave].dropna().values
        bp = ax.boxplot([ds, dn], labels=['Saud.', 'Não\nsaud.'],
                        patch_artist=True, widths=0.5, showmeans=True)
        for patch, cor in zip(bp['boxes'], [COR_SAUD, COR_NAO]):
            patch.set_facecolor(cor)
            patch.set_alpha(0.6)
        ax.set_title(lab, fontweight='bold', fontsize=10)
        ax.grid(alpha=0.3, axis='y')
    fig.suptitle('Intervalos temporais entre pontos fiduciais por grupo',
                 fontweight='bold')
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)


def fig_timeline(df, path):
    ta = np.linspace(-0.3, 0.6, 90)
    grupos = [(df[df['saudavel']], 'Saudável', COR_SAUD),
              (df[~df['saudavel']], 'Não saudável', COR_NAO)]
    fig, axes = plt.subplots(2, 1, figsize=(12, 7.5), sharex=True)
    for ax, (sub, lab, cor) in zip(axes, grupos):
        arr = []
        for e in sub['ensemble']:
            e = np.asarray(e)
            if len(e) != 90:
                continue
            m = np.max(np.abs(e))
            arr.append(e / m if m > 0 else e)
        gm = np.mean(arr, axis=0)
        ax.plot(ta, gm, color=cor, lw=2.2, zorder=1)
        ax.axvline(0, color='gray', ls='--', alpha=0.5)
        ax.axhline(0, color='gray', ls=':', alpha=0.4)
        for nm in PONTOS:
            col = f't_{nm}'
            if col not in sub:
                continue
            mt = np.nanmean(sub[col].values)
            if np.isnan(mt):
                continue
            idx = int(np.argmin(np.abs(ta - mt)))
            rot, cf = FID_INFO[nm]
            ax.scatter([mt], [gm[idx]], c=cf, s=120, marker='o',
                       edgecolors='black', linewidths=0.7, zorder=3)
            ax.annotate(rot, xy=(mt, gm[idx]), xytext=(0, 11),
                        textcoords='offset points', ha='center',
                        fontsize=8.5, fontweight='bold', color=cf, zorder=4)
        ax.set_title(f'Grupo {lab} (n={len(sub)})', fontweight='bold')
        ax.set_ylabel('Amplitude norm.')
        ax.grid(alpha=0.3)
    axes[-1].set_xlabel('Tempo (s)')
    fig.suptitle('Posição média dos pontos fiduciais na onda SCG por grupo',
                 fontweight='bold', fontsize=13)
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)


def escolher_melhor_sinal(df):
    """Seleciona o sinal mais limpo (maior SNR do batimento, FC fisiológica)."""
    melhor, melhor_score = None, -np.inf
    for _, p in df[df['saudavel']].iterrows():
        cache = CACHE_DIR / f"paciente_{p['num']}.npz"
        if not cache.exists():
            continue
        d = np.load(cache)
        cycles = d['cycles']
        if len(cycles) < 20:
            continue
        rr = np.diff(np.sort(d['peaks'])) / FS * 1000.0
        rr = rr[(rr > 300) & (rr < 2000)]
        if len(rr) < 12:
            continue
        fc = 60000.0 / np.mean(rr)
        ens = np.mean(cycles, axis=0)
        std = np.std(cycles, axis=0)
        # SNR: amplitude do batimento vs. dispersão entre ciclos
        snr = np.max(np.abs(ens)) / (np.mean(std) + 1e-9)
        cv = np.std(rr) / np.mean(rr)
        score = snr * (1.0 - min(cv, 0.5))
        if 55 <= fc <= 95 and score > melhor_score:
            melhor_score, melhor = score, int(p['num'])
    return melhor if melhor is not None else int(df['num'].iloc[0])


def fig_sinal_fc_fr(num, path):
    """SCG (FC) e curva respiratória (FR) no mesmo gráfico, com filtros."""
    d = np.load(CACHE_DIR / f"paciente_{num}.npz")
    sinal = np.asarray(d['sinal'])
    fs = FS
    t = np.arange(len(sinal)) / fs

    # --- Filtro cardíaco: passa-banda 2-39 Hz (componente SCG) ---
    bc, ac = signal.butter(4, [2 / (fs / 2), 39 / (fs / 2)], btype='band')
    scg = signal.filtfilt(bc, ac, sinal)

    # --- Batimentos validados (do cache) -> FC ---
    beats = np.sort(np.asarray(d['peaks']))
    rr = np.diff(beats) / fs * 1000.0
    rr = rr[(rr > 300) & (rr < 2000)]
    fc = 60000.0 / np.mean(rr) if len(rr) else np.nan

    # --- Sinal respiratório: envelope de Hilbert -> banda 0,1-0,5 Hz (FR) ---
    env = np.abs(signal.hilbert(scg))
    env0 = env - np.mean(env)
    br, ar = signal.butter(2, [0.1 / (fs / 2), 0.5 / (fs / 2)], btype='band')
    resp = signal.filtfilt(br, ar, env0)
    f, pxx = signal.welch(resp, fs=fs, nperseg=min(len(resp), int(30 * fs)))
    banda_r = (f >= 0.1) & (f <= 0.5)
    fr = f[banda_r][np.argmax(pxx[banda_r])] * 60.0 if np.any(banda_r) else np.nan

    # --- Espectro do ENVELOPE: mostra picos de FR (resp.) e FC (cardíaca) ---
    fe, pe = signal.welch(env0, fs=fs, nperseg=min(len(env0), int(40 * fs)))
    pe_n = pe / (np.max(pe) + 1e-12)

    # Janela de exibição (20 s, para enxergar batimentos e respiração)
    w0, w1 = 10.0, 30.0
    m = (t >= w0) & (t <= w1)
    bmask = beats[(t[beats] >= w0) & (t[beats] <= w1)]

    fig, (ax1, ax3) = plt.subplots(2, 1, figsize=(12, 7),
                                   gridspec_kw={'height_ratios': [2.2, 1]})

    # Painel (a): SCG + respiração no mesmo eixo temporal
    ax1.plot(t[m], scg[m], color='#2c7fb8', lw=0.9, label='SCG filtrado (2–39 Hz)')
    ax1.plot(t[bmask], scg[bmask], 'v', color='#d7301f', ms=7,
             label=f'Batimentos → FC = {fc:.0f} bpm')
    ax1.set_ylabel('Aceleração (m/s²)', color='#2c7fb8')
    ax1.tick_params(axis='y', labelcolor='#2c7fb8')

    ax2 = ax1.twinx()
    resp_disp = resp / (np.max(np.abs(resp[m])) + 1e-9)
    ax2.plot(t[m], resp_disp[m], color='#238b45', lw=2.4,
             label=f'Respiração (0,1–0,5 Hz) → FR = {fr:.0f} rpm')
    ax2.set_ylabel('Sinal respiratório (norm.)', color='#238b45')
    ax2.tick_params(axis='y', labelcolor='#238b45')
    ax2.set_ylim(-1.6, 1.6)

    l1, la1 = ax1.get_legend_handles_labels()
    l2, la2 = ax2.get_legend_handles_labels()
    ax1.legend(l1 + l2, la1 + la2, loc='upper right', fontsize=8.5, ncol=1)
    ax1.set_title(f'(a) Captação simultânea de FC e FR — Paciente #{num} '
                  f'(sinal representativo)', fontweight='bold')
    ax1.set_xlabel('Tempo (s)')
    ax1.grid(alpha=0.25)

    # Painel (b): espectro do envelope evidenciando os picos de FR e FC
    ax3.plot(fe, pe_n, color='#555555', lw=1.4)
    ax3.axvspan(0.1, 0.5, color='#238b45', alpha=0.12, label='Banda respiratória')
    if not np.isnan(fr):
        ax3.axvline(fr / 60.0, color='#238b45', ls='--', lw=1.8,
                    label=f'FR ≈ {fr:.0f} rpm ({fr/60:.2f} Hz)')
    if not np.isnan(fc):
        ax3.axvline(fc / 60.0, color='#d7301f', ls='--', lw=1.8,
                    label=f'FC ≈ {fc:.0f} bpm ({fc/60:.2f} Hz)')
    ax3.set_xlim(0, 3)
    ax3.set_xlabel('Frequência (Hz)')
    ax3.set_ylabel('PSD do envelope (norm.)')
    ax3.set_title('(b) Espectro do envelope: picos respiratório e cardíaco',
                  fontweight='bold')
    ax3.legend(fontsize=8.5, loc='upper right')
    ax3.grid(alpha=0.25)

    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches='tight')
    plt.close(fig)
    return num, fc, fr


# ============================================================
# DOCUMENTO WORD
# ============================================================

def add_tabela_resumo(doc, df):
    saud = df[df['saudavel']]
    nao = df[~df['saudavel']]
    linhas = [
        ('Nº de participantes', str(len(saud)), str(len(nao)), '—'),
        ('Idade (anos)', resumo_grupo(saud['idade']), resumo_grupo(nao['idade']),
         teste_grupos(saud['idade'], nao['idade'])),
        ('Sexo (M / F)',
         f"{np.sum(saud['sexo']=='M')} / {np.sum(saud['sexo']=='F')}",
         f"{np.sum(nao['sexo']=='M')} / {np.sum(nao['sexo']=='F')}", '—'),
        ('FC (bpm)', resumo_grupo(saud['fc']), resumo_grupo(nao['fc']),
         teste_grupos(saud['fc'], nao['fc'])),
        ('SDNN (ms)', resumo_grupo(saud['sdnn']), resumo_grupo(nao['sdnn']),
         teste_grupos(saud['sdnn'], nao['sdnn'])),
        ('RMSSD (ms)', resumo_grupo(saud['rmssd']), resumo_grupo(nao['rmssd']),
         teste_grupos(saud['rmssd'], nao['rmssd'])),
        ('FR (rpm)', resumo_grupo(saud['fr']), resumo_grupo(nao['fr']),
         teste_grupos(saud['fr'], nao['fr'])),
    ]
    tab = doc.add_table(rows=1, cols=4)
    tab.style = 'Light Grid Accent 1'
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tab.rows[0].cells
    for c, txt in zip(hdr, ['Variável', 'Saudável', 'Não saudável', 'p']):
        c.paragraphs[0].add_run(txt).bold = True
    for var, s, ns, p in linhas:
        cells = tab.add_row().cells
        cells[0].text = var
        cells[1].text = s
        cells[2].text = ns
        cells[3].text = p
    return tab


def add_tabela_intervalos(doc, df):
    saud = df[df['saudavel']]
    nao = df[~df['saudavel']]
    tab = doc.add_table(rows=1, cols=4)
    tab.style = 'Light Grid Accent 1'
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tab.rows[0].cells
    for c, txt in zip(hdr, ['Intervalo', 'Saudável', 'Não saudável', 'p']):
        c.paragraphs[0].add_run(txt).bold = True
    for chave, lab, _a, _b in INTERVALOS:
        cells = tab.add_row().cells
        cells[0].text = lab
        cells[1].text = resumo_grupo(saud[chave])
        cells[2].text = resumo_grupo(nao[chave])
        cells[3].text = teste_grupos(saud[chave], nao[chave])
    return tab


def add_tabela_individual(doc, df):
    tab = doc.add_table(rows=1, cols=7)
    tab.style = 'Light List Accent 1'
    hdr = tab.rows[0].cells
    for c, txt in zip(hdr, ['#', 'Grupo', 'Sexo', 'Idade', 'FC', 'SDNN', 'FR']):
        c.paragraphs[0].add_run(txt).bold = True
    for _, p in df.sort_values(['saudavel', 'num'], ascending=[False, True]).iterrows():
        cells = tab.add_row().cells
        cells[0].text = str(p['num'])
        cells[1].text = 'Saud.' if p['saudavel'] else 'Não saud.'
        cells[2].text = str(p['sexo'])
        cells[3].text = str(p['idade'])
        cells[4].text = f"{p['fc']:.0f}" if not np.isnan(p['fc']) else '—'
        cells[5].text = f"{p['sdnn']:.0f}" if not np.isnan(p['sdnn']) else '—'
        cells[6].text = f"{p['fr']:.0f}" if not np.isnan(p['fr']) else '—'


def h(doc, texto, nivel=1):
    doc.add_heading(texto, level=nivel)


def gerar_documento(df, sinal_info):
    saud = df[df['saudavel']]
    nao = df[~df['saudavel']]

    doc = Document()

    # Título
    t = doc.add_heading('Relatório de Resultados — Análise Sismocardiográfica', level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph('Comparação entre grupos Saudável e Não saudável (Projeto PIBIC)')
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    # 1. Introdução
    h(doc, '1. Introdução e Objetivo', 1)
    doc.add_paragraph(
        'Este relatório apresenta a análise comparativa de sinais de sismocardiografia '
        '(SCG) entre participantes classificados como saudáveis e não saudáveis. '
        f'Foram analisados {len(df)} participantes (planilha Pesquisa_PIBIC_sismo), '
        'avaliando-se parâmetros sociodemográficos, frequência cardíaca (FC), '
        'variabilidade da frequência cardíaca (HRV), frequência respiratória (FR) '
        'estimada e a morfologia média da onda SCG de cada grupo.')

    # 2. Metodologia (resumo)
    h(doc, '2. Metodologia (resumo)', 1)
    doc.add_paragraph(
        'Os sinais de aceleração foram filtrados (passa-banda 2–39 Hz) e segmentados em '
        'batimentos a partir da detecção de picos. Para cada participante calculou-se a '
        'onda média (ensemble). A FC e a HRV (SDNN e RMSSD) foram derivadas dos intervalos '
        'entre batimentos; a FR foi estimada pela modulação de amplitude do envelope do '
        'sinal (banda 0,1–0,5 Hz). As comparações entre grupos utilizaram o teste de '
        'Mann-Whitney (valores expressos como média ± desvio-padrão).')

    # 3. Sociodemográfico
    h(doc, '3. Dados Sociodemográficos', 1)
    doc.add_paragraph(
        f'O grupo saudável reuniu {len(saud)} participantes '
        f'(idade {resumo_grupo(saud["idade"])} anos) e o grupo não saudável '
        f'{len(nao)} participantes (idade {resumo_grupo(nao["idade"])} anos).')
    doc.add_picture(str(FIG_DIR / 'sociodemografico.png'), width=Inches(6.3))
    _legenda(doc, 'Figura 1. Distribuição de participantes, sexo e idade por grupo.')

    # 4. Tabela resumo
    h(doc, '4. Síntese Comparativa das Variáveis', 1)
    add_tabela_resumo(doc, df)
    _legenda(doc, 'Tabela 1. Comparação das variáveis entre grupos (média ± DP; p de Mann-Whitney).')

    # 5. FC + HRV + FR
    h(doc, '5. Frequência Cardíaca, HRV e Frequência Respiratória', 1)
    doc.add_paragraph(
        f'A FC média foi de {resumo_grupo(saud["fc"])} bpm no grupo saudável e '
        f'{resumo_grupo(nao["fc"])} bpm no grupo não saudável. A HRV '
        f'(SDNN: {resumo_grupo(saud["sdnn"])} vs {resumo_grupo(nao["sdnn"])} ms; '
        f'RMSSD: {resumo_grupo(saud["rmssd"])} vs {resumo_grupo(nao["rmssd"])} ms) e a FR '
        f'({resumo_grupo(saud["fr"])} vs {resumo_grupo(nao["fr"])} rpm) são comparadas na Figura 2.')
    doc.add_picture(str(FIG_DIR / 'boxplots.png'), width=Inches(6.5))
    _legenda(doc, 'Figura 2. Distribuição de FC, SDNN, RMSSD e FR por grupo (▲ = média).')

    # 6. Morfologia
    h(doc, '6. Morfologia Média da Onda SCG', 1)
    doc.add_paragraph(
        'A morfologia média (ensemble) de cada grupo é apresentada na Figura 3, tanto na '
        'forma normalizada (para comparar o padrão temporal) quanto em amplitude absoluta '
        '(para comparar a intensidade das vibrações cardíacas).')
    doc.add_picture(str(FIG_DIR / 'morfologia.png'), width=Inches(6.5))
    _legenda(doc, 'Figura 3. Morfologia média da onda SCG por grupo (banda = ± 1 DP).')

    # 7. Intervalos entre pontos fiduciais
    h(doc, '7. Intervalos Temporais entre Pontos Fiduciais', 1)
    doc.add_paragraph(
        'A partir dos 8 pontos fiduciais (Sørensen et al., 2018) detectados na onda média '
        'de cada participante, foram calculados intervalos temporais com significado '
        'fisiológico: contração isovolumétrica (CIV: MC→AO), tempo de ejeção do ventrículo '
        'esquerdo (LVET: AO→AC), relaxamento isovolumétrico (RIV: AC→MO) e início do '
        'enchimento diastólico (MO→pico da onda E). A Tabela 2 resume a comparação entre '
        'grupos e as Figuras 4 e 5 apresentam a distribuição dos intervalos e a posição '
        'média dos pontos sobre a onda SCG.')
    add_tabela_intervalos(doc, df)
    _legenda(doc, 'Tabela 2. Intervalos temporais entre pontos fiduciais '
                  '(média ± DP; p de Mann-Whitney).')
    doc.add_picture(str(FIG_DIR / 'intervalos.png'), width=Inches(6.5))
    _legenda(doc, 'Figura 4. Distribuição dos intervalos temporais por grupo (▲ = média).')
    doc.add_picture(str(FIG_DIR / 'timeline.png'), width=Inches(6.3))
    _legenda(doc, 'Figura 5. Posição média dos pontos fiduciais na onda SCG de cada grupo.')

    # 8. Captação simultânea de FC e FR
    s_num, s_fc, s_fr = sinal_info
    h(doc, '8. Captação Simultânea de FC e FR (Sinal Representativo)', 1)
    doc.add_paragraph(
        'A Figura 6 demonstra, em um sinal de boa qualidade (paciente #%d), a captação '
        'simultânea da frequência cardíaca e da frequência respiratória a partir de um '
        'único acelerômetro. O componente cardíaco é obtido por filtro passa-banda '
        '2–39 Hz (batimentos → FC ≈ %.0f bpm), enquanto o sinal respiratório é extraído '
        'da modulação de amplitude do envelope na banda 0,1–0,5 Hz (→ FR ≈ %.0f rpm). '
        'O espectro (painel b) evidencia os dois picos em bandas de frequência distintas, '
        'confirmando a separação entre as componentes cardíaca e respiratória.'
        % (s_num, s_fc, s_fr))
    doc.add_picture(str(FIG_DIR / 'sinal_fc_fr.png'), width=Inches(6.5))
    _legenda(doc, 'Figura 6. Sinal SCG representativo com captação simultânea de FC '
                  '(batimentos, 2–39 Hz) e FR (respiração, 0,1–0,5 Hz).')

    # 9. Dashboards
    h(doc, '9. Dashboards Individuais por Grupo', 1)
    doc.add_paragraph('Onda média de cada participante, organizada por grupo clínico.')
    doc.add_picture(str(FIG_DIR / 'dashboard_saudavel.png'), width=Inches(6.5))
    _legenda(doc, 'Figura 7. Ondas médias — grupo Saudável.')
    doc.add_picture(str(FIG_DIR / 'dashboard_nao.png'), width=Inches(6.5))
    _legenda(doc, 'Figura 8. Ondas médias — grupo Não saudável.')

    # 10. Apêndice
    doc.add_page_break()
    h(doc, '10. Apêndice — Dados Individuais', 1)
    add_tabela_individual(doc, df)

    out = DATA_DIR / 'Relatorio_PIBIC_SCG.docx'
    doc.save(out)
    return out


def _legenda(doc, texto):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.italic = True
    r.font.size = Pt(9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 60)
    print("GERANDO RELATÓRIO WORD - PIBIC SCG")
    print("=" * 60)
    FIG_DIR.mkdir(exist_ok=True)

    print("Carregando dados e calculando métricas...")
    df = carregar_dados()
    saud = df[df['saudavel']]
    nao = df[~df['saudavel']]
    print(f"  Total: {len(df)} | Saudáveis: {len(saud)} | Não saudáveis: {len(nao)}")

    print("Gerando figuras...")
    fig_sociodemografico(df, FIG_DIR / 'sociodemografico.png')
    fig_morfologia(df, FIG_DIR / 'morfologia.png')
    fig_boxplots(df, FIG_DIR / 'boxplots.png')
    fig_intervalos(df, FIG_DIR / 'intervalos.png')
    fig_timeline(df, FIG_DIR / 'timeline.png')
    fig_dashboard(saud, 'Grupo Saudável — Onda Média SCG', COR_SAUD,
                  FIG_DIR / 'dashboard_saudavel.png')
    fig_dashboard(nao, 'Grupo Não Saudável — Onda Média SCG', COR_NAO,
                  FIG_DIR / 'dashboard_nao.png')

    print("Selecionando sinal representativo e gerando figura FC/FR...")
    melhor = escolher_melhor_sinal(df)
    sinal_info = fig_sinal_fc_fr(melhor, FIG_DIR / 'sinal_fc_fr.png')
    print(f"  Sinal representativo: paciente #{sinal_info[0]} "
          f"(FC={sinal_info[1]:.0f} bpm, FR={sinal_info[2]:.0f} rpm)")

    print("Montando documento...")
    out = gerar_documento(df, sinal_info)

    print("\n" + "=" * 60)
    print(f"RELATÓRIO GERADO: {out.name}")
    print(f"  Figuras em: {FIG_DIR}")
    print("=" * 60)


if __name__ == '__main__':
    main()
